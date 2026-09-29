import sys
import os
import argparse

def render_pdf(input_path, output_path):
    import pymupdf
    doc = pymupdf.open(input_path)
    if len(doc) == 0:
        raise ValueError("PDF has no pages")
    page = doc[0]
    pix = page.get_pixmap(dpi=150)
    pix.save(output_path)
    return True

def render_excel(input_path, output_path):
    # Try Excel COM first if on Windows
    try:
        import win32com.client
        import pymupdf
        import tempfile

        temp_pdf = os.path.join(tempfile.gettempdir(), f"excel_temp_{os.getpid()}.pdf")
        xl = win32com.client.Dispatch("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        try:
            wb = xl.Workbooks.Open(os.path.abspath(input_path))
            wb.ExportAsFixedFormat(0, temp_pdf)
            wb.Close(False)
        finally:
            xl.Quit()

        if os.path.exists(temp_pdf):
            render_pdf(temp_pdf, output_path)
            try:
                os.remove(temp_pdf)
            except Exception:
                pass
            return True
    except Exception as e:
        sys.stderr.write(f"Excel COM warning: {e}. Falling back to openpyxl + Pillow render.\n")

    # Fallback to openpyxl + Pillow
    import openpyxl
    from PIL import Image, ImageDraw, ImageFont

    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.active

    # Determine bounds
    max_r = min(35, ws.max_row or 1)
    max_c = min(10, ws.max_column or 1)

    # Image canvas setup
    img_w = 1200
    img_h = 1500
    img = Image.new("RGB", (img_w, img_h), color="#FFFFFF")
    draw = ImageDraw.Draw(img)

    # Top Excel-like green bar
    draw.rectangle([0, 0, img_w, 48], fill="#107C41")
    filename = os.path.basename(input_path)
    sheet_title = getattr(ws, 'title', 'Hoja 1')

    # Draw header text
    try:
        font_title = ImageFont.truetype("arial.ttf", 22)
        font_cell = ImageFont.truetype("arial.ttf", 14)
        font_grid = ImageFont.truetype("arial.ttf", 12)
    except Exception:
        font_title = ImageFont.load_default()
        font_cell = ImageFont.load_default()
        font_grid = ImageFont.load_default()

    draw.text((20, 12), f"Excel: {filename} - [{sheet_title}]", fill="#FFFFFF", font=font_title)

    # Column and row coordinates
    col_w = max(90, (img_w - 70) // max_c)
    row_h = 32
    start_y = 70
    start_x = 50

    # Draw column header (A, B, C...)
    draw.rectangle([0, start_y, img_w, start_y + row_h], fill="#F3F2F1")
    for c in range(1, max_c + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        cx = start_x + (c - 1) * col_w
        draw.rectangle([cx, start_y, cx + col_w, start_y + row_h], outline="#D2D0CE", width=1)
        draw.text((cx + col_w // 2 - 5, start_y + 8), col_letter, fill="#605E5C", font=font_grid)

    # Draw rows and cells
    for r in range(1, max_r + 1):
        ry = start_y + r * row_h
        if ry + row_h > img_h:
            break
        # Row number label
        draw.rectangle([0, ry, start_x, ry + row_h], fill="#F3F2F1", outline="#D2D0CE", width=1)
        draw.text((15, ry + 8), str(r), fill="#605E5C", font=font_grid)

        # Cells in row
        for c in range(1, max_c + 1):
            cx = start_x + (c - 1) * col_w
            draw.rectangle([cx, ry, cx + col_w, ry + row_h], outline="#E1DFDD", width=1)
            cell_val = ws.cell(row=r, column=c).value
            if cell_val is not None:
                val_str = str(cell_val).strip()
                if len(val_str) > 18:
                    val_str = val_str[:16] + ".."
                draw.text((cx + 6, ry + 8), val_str, fill="#201F1E", font=font_cell)

    img.save(output_path, "PNG")
    return True

def render_word(input_path, output_path):
    try:
        import win32com.client
        import pymupdf
        import tempfile

        temp_pdf = os.path.join(tempfile.gettempdir(), f"word_temp_{os.getpid()}.pdf")
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        try:
            doc = word.Documents.Open(os.path.abspath(input_path))
            doc.SaveAs2(temp_pdf, FileFormat=17)
            doc.Close()
        finally:
            word.Quit()

        if os.path.exists(temp_pdf):
            render_pdf(temp_pdf, output_path)
            try:
                os.remove(temp_pdf)
            except Exception:
                pass
            return True
    except Exception as e:
        sys.stderr.write(f"Word COM error: {e}\n")
        raise

def main():
    parser = argparse.ArgumentParser(description="Render first page screenshot of PDF, Excel, Word")
    parser.add_argument("--input", required=True, help="Input file path")
    parser.add_argument("--output", required=True, help="Output PNG path")
    args = parser.parse_args()

    ext = os.path.splitext(args.input)[1].lower()
    if ext == ".pdf":
        render_pdf(args.input, args.output)
    elif ext in [".xlsx", ".xls"]:
        render_excel(args.input, args.output)
    elif ext in [".docx", ".doc"]:
        render_word(args.input, args.output)
    else:
        raise ValueError(f"Unsupported format for screenshot: {ext}")

    print(f"SUCCESS: {args.output}")

if __name__ == "__main__":
    main()
